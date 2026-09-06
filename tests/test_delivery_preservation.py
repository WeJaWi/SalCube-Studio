"""Preserve owner-added thumbnails and OOXML content during title updates."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image

from pipeline import delivery_doc


def fixture(path):
    picture = io.BytesIO()
    Image.new("RGB", (320, 180), (25, 100, 145)).save(picture, "PNG")
    picture.seek(0)
    doc = Document()
    p = doc.add_paragraph("Existing Video Title")
    p.runs[0].bold = True
    doc.add_picture(picture, width=Inches(3.5))
    doc.add_paragraph("Owner notes below the thumbnail")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Title Inside a Table"
    doc.sections[0].header.paragraphs[0].text = "Owner header"
    doc.sections[0].footer.paragraphs[0].text = "Owner footer"
    doc.save(path)
    with ZipFile(path, "a") as package:
        package.writestr("customXml/owner-notes.xml", b"<notes>Keep this extension</notes>")


class DeliveryPreservationTests(unittest.TestCase):
    def test_append_keeps_thumbnail_and_every_other_package_part(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Titles.docx"
            fixture(path)
            with ZipFile(path) as package:
                before = {name: package.read(name) for name in package.namelist()}
            self.assertTrue(delivery_doc.append_titles(["New Video Title"], path))
            with ZipFile(path) as package:
                self.assertEqual(set(before), set(package.namelist()))
                for name, data in before.items():
                    if name != "word/document.xml":
                        self.assertEqual(package.read(name), data, name)
                after_xml = package.read("word/document.xml")
            old_body = etree.fromstring(before["word/document.xml"]).find(qn("w:body"))
            new_body = etree.fromstring(after_xml).find(qn("w:body"))
            # Existing paragraph/drawing/table XML is unchanged, in order.
            for old, new in zip(list(old_body)[:-1], list(new_body)[:-2]):
                self.assertEqual(etree.tostring(old), etree.tostring(new))
            updated = Document(path)
            self.assertEqual(len(updated.inline_shapes), 1)
            self.assertEqual(updated.paragraphs[-1].text, "New Video Title")
            self.assertTrue(updated.paragraphs[0].runs[0].bold)
            self.assertEqual(updated.sections[0].header.paragraphs[0].text, "Owner header")

    def test_duplicate_title_is_byte_identical_and_table_titles_are_seen(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Titles.docx"
            fixture(path)
            before = path.read_bytes()
            self.assertFalse(delivery_doc.append_titles(["existing video title", "Title Inside a Table"], path))
            self.assertEqual(path.read_bytes(), before)

    def test_invalid_document_is_retained(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Titles.docx"
            path.write_bytes(b"not a valid Word file")
            before = path.read_bytes()
            with self.assertRaises(Exception):
                delivery_doc.append_titles(["New title"], path)
            self.assertEqual(path.read_bytes(), before)

    def test_failed_replace_leaves_original_and_no_partial_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Titles.docx"
            fixture(path)
            before = path.read_bytes()
            with patch.object(delivery_doc.os, "replace", side_effect=OSError("disk error")):
                with self.assertRaises(OSError):
                    delivery_doc.append_titles(["New title"], path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(Path(tmp).glob(".docx-append-*")), [])


class ProfileDocumentUploadTests(unittest.TestCase):
    def setUp(self):
        import app
        from tests.isolation import isolated_runtime
        self.app = app
        self.context = isolated_runtime()
        self.root = self.context.__enter__()
        self.addCleanup(self.context.__exit__, None, None, None)
        self.remote_path = self.root / "remote.docx"
        fixture(self.remote_path)
        self.original = self.remote_path.read_bytes()

    def upload(self, title, *, concurrent_edit=False, corrupt=False):
        from unittest.mock import Mock
        app = self.app
        remote = {"id": "doc-id", "name": delivery_doc.DOCX_NAME, "modifiedTime": "version-1"}
        drive = Mock(configured=True, authorized=True)
        drive.ensure_profile_folder.return_value = "profile-folder"
        drive.find_file.side_effect = [remote, {**remote, "modifiedTime": "version-2"} if concurrent_edit else remote]
        drive.download_to_path.side_effect = lambda _id, path: path.write_bytes(
            b"invalid Word file" if corrupt else self.original)
        drive.upload.return_value = remote
        app.JOBS["doc-job"] = {"profile": "cars", "status": "done", "title": title, "log": []}
        with patch.object(app, "DRIVE", drive), patch.object(app, "_persist_queue"), \
                patch.object(app.traceback, "print_exc"):
            app._upload_profile_docx("doc-job", self.root / "doc-job")
        return drive, app.config.OUTPUT_DIR / "profile-documents" / "cars" / delivery_doc.DOCX_NAME

    def test_upload_preserves_remote_thumbnail_and_saves_recovery_copy(self):
        drive, local = self.upload("New Completed Video")
        self.assertEqual(len(Document(local).inline_shapes), 1)
        self.assertEqual(Document(local).paragraphs[-1].text, "New Completed Video")
        self.assertEqual(local.with_suffix(".previous.docx").read_bytes(), self.original)
        drive.upload.assert_called_once()
        self.assertEqual(drive.upload.call_args.kwargs["file_id"], "doc-id")
        self.assertEqual(self.app.JOBS["doc-job"]["package"]["state"], "uploaded")

    def test_existing_title_does_not_overwrite_remote_document(self):
        drive, local = self.upload("Existing Video Title")
        drive.upload.assert_not_called()
        self.assertEqual(local.read_bytes(), self.original)

    def test_detected_owner_edit_is_not_overwritten(self):
        drive, _ = self.upload("New Completed Video", concurrent_edit=True)
        drive.upload.assert_not_called()
        self.assertEqual(self.app.JOBS["doc-job"]["package"]["state"], "error")
        self.assertIn("changed during this update", self.app.JOBS["doc-job"]["package"]["error"])

    def test_invalid_remote_document_is_not_replaced_with_title_only(self):
        drive, _ = self.upload("New Completed Video", corrupt=True)
        drive.upload.assert_not_called()
        self.assertEqual(self.app.JOBS["doc-job"]["package"]["state"], "error")


if __name__ == "__main__":
    unittest.main()
