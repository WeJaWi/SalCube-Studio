"""Append titles to a profile document without discarding owner-added content."""
from __future__ import annotations

from pathlib import Path
from typing import Iterable
from copy import copy
import os
import tempfile
from zipfile import ZipFile

from lxml import etree

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX_NAME = "Titles and Thumbnails.docx"


def _style_font(style, name: str, size: float, color: str = "111111",
                *, bold: bool = False) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def _configure_styles(doc: Document) -> None:
    """Resolve the compact_reference_guide preset into Word styles."""
    normal = doc.styles["Normal"]
    _style_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        _style_font(style, "Calibri", size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def normalize_titles(titles: Iterable[str]) -> list[str]:
    """Trim titles and preserve their first-seen order without duplicates."""
    if isinstance(titles, str):
        titles = [titles]
    clean: list[str] = []
    seen: set[str] = set()
    for value in titles:
        title = " ".join(str(value or "").split()).strip()
        key = title.casefold()
        if title and key not in seen:
            clean.append(title)
            seen.add(key)
    return clean


def read_titles(path: Path) -> list[str]:
    """Read the title paragraphs from an existing cumulative document."""
    path = Path(path)
    if not path.is_file():
        return []
    return normalize_titles(p.text for p in Document(path).paragraphs)


def create(titles: Iterable[str], dest: Path) -> Path:
    """Write one plain paragraph per title and no thumbnail images or metadata."""
    clean = normalize_titles(titles)
    if not clean:
        raise ValueError("at least one video title is required")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(doc)

    # The owner asked for titles only: no document heading, dates, numbering,
    # job ids, tables, or embedded thumbnail images.
    for title in clean:
        paragraph = doc.add_paragraph(title)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = True

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest


def append_titles(titles: Iterable[str], dest: Path) -> bool:
    """Append only missing titles, preserving every existing package part.

    Rebuilding from paragraph text drops drawings. A general document
    round-trip may also rewrite unsupported Word extensions. Update only
    document.xml and copy images, relationships, comments, styles, headers,
    embedded objects, and all other entries byte-for-byte.
    """
    dest = Path(dest)
    clean = normalize_titles(titles)
    if not clean:
        raise ValueError("at least one video title is required")
    if not dest.exists():
        create(clean, dest)
        return True
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(dest) as package:
        original = package.read("word/document.xml")
        document = etree.fromstring(original, parser)
        body = document.find(qn("w:body"))
        if body is None:
            raise ValueError("Word document has no body; original retained")
        existing = normalize_titles("".join(p.itertext(qn("w:t")))
                                    for p in body.iter(qn("w:p")))
        seen = {title.casefold() for title in existing}
        additions = [title for title in clean if title.casefold() not in seen]
        if not additions:
            return False  # no serialization or upload on an idempotent retry
        section = body.find(qn("w:sectPr"))
        for title in additions:
            paragraph = etree.Element(qn("w:p"))
            props = etree.SubElement(paragraph, qn("w:pPr"))
            etree.SubElement(props, qn("w:keepLines"))
            etree.SubElement(props, qn("w:widowControl"))
            run = etree.SubElement(paragraph, qn("w:r"))
            etree.SubElement(run, qn("w:t")).text = title
            if section is None:
                body.append(paragraph)
            else:
                section.addprevious(paragraph)
        updated = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone=True)
        fd, name = tempfile.mkstemp(prefix=".docx-append-", suffix=".docx", dir=dest.parent)
        os.close(fd)
        temp = Path(name)
        try:
            with ZipFile(temp, "w") as result:
                result.comment = package.comment
                for info in package.infolist():
                    # Writing mutates ZipInfo offsets. Keep the source metadata
                    # intact so verification still reads the original package.
                    result.writestr(copy(info), updated if info.filename == "word/document.xml"
                                    else package.read(info.filename))
            with ZipFile(temp) as result:
                if result.testzip() is not None:
                    raise ValueError("Updated Word document failed validation; original retained")
                for info in package.infolist():
                    if info.filename != "word/document.xml" and result.read(info.filename) != package.read(info.filename):
                        raise ValueError("Existing Word content changed unexpectedly; original retained")
            # Close the input before replace so this is also valid on Windows.
        except BaseException:
            temp.unlink(missing_ok=True)
            raise
    try:
        os.replace(temp, dest)
    finally:
        temp.unlink(missing_ok=True)
    return True
